from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from .common import repo_path, utc_now_iso
from .detail_maps import DETAIL_MAP_SECTIONS, slugify


SUBJECT_CARD_CANDIDATES_HEADER = "candidate_id,title,proposed_slug,source,discovery_basis,linked_features,linked_detail_maps,primary_objects,subject_type,confidence,proposed_action,status,notes"
SUBJECT_CARD_CONTOURS_HEADER = "contour_id,slug,title,contour_type,linked_features,primary_objects,linked_detail_maps,scenario_summary,migration_boundary,why_this_is_one_contour,why_not_technical_bucket,evidence_refs,runtime_refs,technical_bucket_refs,status,confidence,notes"
SUBJECT_CARD_CLASSIFICATION_HEADER = "candidate_id,proposed_slug,decision,subject_type,registry_slug,merge_into,split_from,why_separate_card,status,confidence,notes"
SUBJECT_CARD_REGISTRY_HEADER = "slug,title,subject_type,status,confidence,origin_layer,owner_feature,linked_features,linked_detail_maps,primary_objects,coverage_scope,why_separate_card,merge_into,split_from,card_path,evidence_count,gap_count,review_notes"
SUBJECT_CARD_COVERAGE_HEADER = "source_kind,source_id,feature_id,detail_map_slug,subject_card_slug,relation,confidence,notes"
SUBJECT_CARD_EVIDENCE_HEADER = "evidence_id,section,claim,source_type,source_path,line,linked_diff_id,linked_feature_id,confidence,notes"
SUBJECT_CARD_GAPS_HEADER = "gap_id,section,question,needed_source,status,blocking,notes"
SUBJECT_CARD_TYPES = {
    "business_process",
    "business_document",
    "reference_model",
    "integration",
    "access_model",
    "ui_surface",
    "background_automation",
    "technical_support",
}
SUBJECT_CARD_ACTIONS = {"accept", "split", "merge", "reject", "supporting"}
SUBJECT_CARD_STATUSES = {
    "candidate",
    "accepted",
    "draft",
    "needs_static_analysis",
    "needs_runtime_data",
    "needs_ui_check",
    "needs_review",
    "ready_for_review",
    "reviewed",
    "rejected",
    "merged_into_other",
    "supporting",
    "unclassified",
}
SUBJECT_CARD_SECTIONS = tuple(DETAIL_MAP_SECTIONS)
GENERIC_BUCKET_SUFFIXES = (
    "документы и журналы",
    "регистры и движения",
    "общие модули и платформенная логика",
    "формы, команды и интерфейс",
    "прочие связанные объекты",
)
GENERIC_TEMPLATE_FRAGMENTS = (
    "имеет собственные объекты, формы или источники риска",
    "подтверждена статическим clean diff",
    "подтверждено, что `",
    "карточка готова к первичному аналитическому ревью",
)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def non_empty_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return [row for row in rows if any((value or "").strip() for value in row.values())]


def write_csv_rows(path: Path, header: str, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = header.split(",")
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def subject_root(root: Path) -> Path:
    return repo_path(root, "analysis/subject-cards")


def ensure_scaffold(root: Path) -> None:
    base = subject_root(root)
    (base / "_templates").mkdir(parents=True, exist_ok=True)
    (base / "cards").mkdir(parents=True, exist_ok=True)
    readme = base / "README.md"
    if not readme.exists():
        readme.write_text(
            "# Предметные карточки доработок\n\n"
            "Этот каталог хранит итерационный слой `subject cards`: кандидаты, реестр, карточки, evidence и gaps.\n\n"
            "Базовый цикл:\n\n"
            "```bash\n"
            "python -m one_c_autoresearch subject-card discover\n"
            "python -m one_c_autoresearch subject-card contour-draft\n"
            "python -m one_c_autoresearch subject-card contour-validate\n"
            "python -m one_c_autoresearch subject-card classify\n"
            "python -m one_c_autoresearch subject-card registry-build\n"
            "python -m one_c_autoresearch subject-card seed --from-registry\n"
            "python -m one_c_autoresearch subject-card refine --card <slug>\n"
            "python -m one_c_autoresearch subject-card validate\n"
            "python -m one_c_autoresearch review-dashboard build\n"
            "```\n",
            encoding="utf-8",
            newline="\n",
        )
    template = base / "_templates/subject-card.json"
    if not template.exists():
        template.write_text(
            json.dumps(
                {
                    "schema_version": "subject-card/v1",
                    "slug": "example-subject",
                    "title": "Пример предметной доработки",
                    "status": "draft",
                    "confidence": "medium",
                    "subject_type": "business_process",
                    "origin_layer": "manual",
                    "primary_objects": [],
                    "coverage_scope": "Пока не определено.",
                    "why_separate_card": "Пока не определено.",
                    "merge_into": "",
                    "split_from": "",
                    "identification": "Как предметная доработка определяется в исходниках или данных ИБ.",
                    "summary": "Краткое описание предметной доработки.",
                    "key_conclusion": "Главный вывод для аналитика.",
                    "upgrade_risk": "Риск перехода на целевой релиз.",
                    "linked_features": [],
                    "linked_detail_maps": [],
                    "source_mode": "template",
                    "source_artifacts": [],
                    "sections": {section: [] for section in SUBJECT_CARD_SECTIONS},
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
            newline="\n",
        )
    for relative, header in (
        ("candidates.csv", SUBJECT_CARD_CANDIDATES_HEADER),
        ("contours.csv", SUBJECT_CARD_CONTOURS_HEADER),
        ("classification.csv", SUBJECT_CARD_CLASSIFICATION_HEADER),
        ("registry.csv", SUBJECT_CARD_REGISTRY_HEADER),
        ("coverage.csv", SUBJECT_CARD_COVERAGE_HEADER),
    ):
        path = base / relative
        if not path.exists():
            write_csv_rows(path, header, [])


def split_refs(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [part.strip() for part in str(value or "").replace("\n", ";").split(";") if part.strip()]


def join_refs(values: list[str]) -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        item = str(value or "").strip()
        if not item or item in seen:
            continue
        seen.add(item)
        result.append(item)
    return ";".join(result)


def first_sentence(text: str, limit: int = 280) -> str:
    compact = " ".join(str(text or "").split())
    if not compact:
        return ""
    for marker in (". ", "; "):
        idx = compact.find(marker)
        if 40 <= idx <= limit:
            return compact[: idx + 1]
    return compact[:limit].rstrip()


def generic_bucket_title(title: str) -> bool:
    lowered = str(title or "").strip().lower()
    return any(lowered.endswith(suffix) for suffix in GENERIC_BUCKET_SUFFIXES)


def generic_template_text(text: str) -> bool:
    lowered = str(text or "").lower()
    return any(fragment in lowered for fragment in GENERIC_TEMPLATE_FRAGMENTS)


def spreadsheet_reference(value: str) -> bool:
    lowered = value.lower()
    return any(suffix in lowered for suffix in (".xlsx", ".xlsm", ".xlsb", ".xls#"))


def read_detail_maps(root: Path) -> list[tuple[Path, dict[str, Any]]]:
    detail_root = repo_path(root, "analysis/detail-maps")
    if not detail_root.exists():
        return []
    maps: list[tuple[Path, dict[str, Any]]] = []
    for path in sorted(detail_root.rglob("detail-map.json")):
        relative_parts = path.relative_to(detail_root).parts
        if "_templates" in relative_parts:
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        maps.append((path, data))
    return maps


def accepted_detail_maps(root: Path) -> list[tuple[Path, dict[str, Any]]]:
    result: list[tuple[Path, dict[str, Any]]] = []
    detail_root = repo_path(root, "analysis/detail-maps")
    generated_root = detail_root / "generated"
    for path, data in read_detail_maps(root):
        try:
            path.relative_to(generated_root)
            continue
        except ValueError:
            pass
        if str(data.get("generation_mode") or "").strip() in {"manual", "enriched"}:
            result.append((path, data))
    return result


def discover_candidates(root: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for path, data in accepted_detail_maps(root):
        slug = str(data.get("slug") or path.parent.name).strip()
        if not slug or slug in seen:
            continue
        seen.add(slug)
        rows.append(
            {
                "candidate_id": f"SC-{len(rows) + 1:04d}",
                "title": str(data.get("title") or slug),
                "proposed_slug": slug,
                "source": path.relative_to(root).as_posix(),
                "discovery_basis": "Существующая аналитическая detail-map вне generated/ подходит как гипотеза предметной карточки.",
                "linked_features": join_refs(split_refs(data.get("linked_features", []))),
                "linked_detail_maps": slug,
                "primary_objects": join_refs(primary_objects_from_detail_map(data)),
                "subject_type": subject_type_for_detail_map(data),
                "confidence": str(data.get("confidence") or "medium"),
                "proposed_action": "accept",
                "status": "accepted",
                "notes": "Registry-build сохраняет решение отдельно от discover; seed переносит факты только из утвержденного registry.",
            }
        )

    feature_rows = non_empty_rows(read_csv_rows(repo_path(root, "analysis/indexes/final-feature-map.csv")))
    if not feature_rows:
        feature_rows = non_empty_rows(read_csv_rows(repo_path(root, "analysis/indexes/feature-map.csv")))
    details_by_feature = detail_slugs_by_feature(root)
    for row in feature_rows:
        feature_id = (row.get("feature_id") or "").strip()
        title = (row.get("title") or feature_id).strip()
        slug = slugify(f"{feature_id}-{title}") if feature_id else slugify(title)
        if not slug or slug in seen:
            continue
        seen.add(slug)
        subject_type = subject_type_for_feature(row)
        proposed_action = "split"
        status = "candidate"
        rows.append(
            {
                "candidate_id": f"SC-{len(rows) + 1:04d}",
                "title": title,
                "proposed_slug": slug,
                "source": "analysis/indexes/final-feature-map.csv",
                "discovery_basis": "BF-блок является контейнером фактов и гипотезой для дальнейшей классификации в предметные карточки.",
                "linked_features": feature_id,
                "linked_detail_maps": join_refs(details_by_feature.get(feature_id, [])),
                "primary_objects": row.get("source_bucket", ""),
                "subject_type": subject_type,
                "confidence": row.get("confidence", "") or "medium",
                "proposed_action": proposed_action,
                "status": status,
                "notes": "Не считать готовой предметной карточкой по правилу один BF = одна карточка; нужен classify/registry-build.",
            }
        )
    return rows


def discover_subject_cards(root: Path) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    rows = discover_candidates(root)
    write_csv_rows(subject_root(root) / "candidates.csv", SUBJECT_CARD_CANDIDATES_HEADER, rows)
    return {"candidates": len(rows), "accepted": sum(1 for row in rows if row["status"] == "accepted")}


def empty_sections() -> dict[str, list[dict[str, str]]]:
    return {section: [] for section in SUBJECT_CARD_SECTIONS}


def normalize_section_rows(rows: Any) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    if not isinstance(rows, list):
        return result
    for row in rows:
        if isinstance(row, dict):
            result.append({str(key): "" if value is None else str(value) for key, value in row.items()})
    return result


def detail_map_by_slug(root: Path) -> dict[str, tuple[Path, dict[str, Any]]]:
    result: dict[str, tuple[Path, dict[str, Any]]] = {}
    for path, data in read_detail_maps(root):
        slug = str(data.get("slug") or path.parent.name).strip()
        if slug:
            result[slug] = (path, data)
    return result


def feature_by_id(root: Path) -> dict[str, dict[str, str]]:
    rows = non_empty_rows(read_csv_rows(repo_path(root, "analysis/indexes/final-feature-map.csv")))
    if not rows:
        rows = non_empty_rows(read_csv_rows(repo_path(root, "analysis/indexes/feature-map.csv")))
    return {(row.get("feature_id") or "").strip(): row for row in rows if (row.get("feature_id") or "").strip()}


def scenario_artifacts_by_feature(root: Path) -> dict[str, dict[str, Path]]:
    result: dict[str, dict[str, Path]] = {}
    scenario_root = repo_path(root, "analysis/reverse-map/scenarios")
    if not scenario_root.exists():
        return result
    for path in sorted(scenario_root.iterdir()):
        if not path.is_dir():
            continue
        feature_id = path.name.upper()
        artifacts: dict[str, Path] = {}
        for name in ("summary.md", "evidence.csv", "decisions.csv"):
            artifact = path / name
            if artifact.exists():
                artifacts[name] = artifact
        for artifact in sorted(path.glob("runtime-*.csv")):
            artifacts[artifact.name] = artifact
        if artifacts:
            result[feature_id] = artifacts
    return result


def infobase_refs_by_feature(root: Path) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    checks = non_empty_rows(read_csv_rows(repo_path(root, "analysis/reverse-map/infobase-checks.csv")))
    for row in checks:
        feature_id = (row.get("feature_id") or "").strip()
        check_id = (row.get("check_id") or "").strip()
        if not feature_id or not check_id:
            continue
        if row.get("status_after_pass") == "closed":
            result.setdefault(feature_id, []).append(f"analysis/reverse-map/infobase-checks.csv#{check_id}")
    return {key: split_refs(join_refs(value)) for key, value in result.items()}


def contour_rows(root: Path) -> list[dict[str, str]]:
    return non_empty_rows(read_csv_rows(subject_root(root) / "contours.csv"))


def accepted_contours(root: Path) -> dict[str, dict[str, str]]:
    return {row.get("slug", ""): row for row in contour_rows(root) if row.get("slug") and row.get("status") == "accepted"}


def contour_by_feature(root: Path) -> dict[str, list[dict[str, str]]]:
    result: dict[str, list[dict[str, str]]] = {}
    for row in contour_rows(root):
        if row.get("status") != "accepted":
            continue
        for feature_id in split_refs(row.get("linked_features", "")):
            result.setdefault(feature_id, []).append(row)
    return result


def candidate_by_slug(root: Path) -> dict[str, dict[str, str]]:
    return {row.get("proposed_slug", ""): row for row in candidate_rows(root) if row.get("proposed_slug")}


def draft_subject_card_contours(root: Path) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    features = feature_by_id(root)
    candidates = candidate_rows(root)
    candidates_by_feature: dict[str, list[dict[str, str]]] = {}
    for candidate in candidates:
        for feature_id in split_refs(candidate.get("linked_features", "")):
            candidates_by_feature.setdefault(feature_id, []).append(candidate)
    scenario_artifacts = scenario_artifacts_by_feature(root)
    runtime_refs = infobase_refs_by_feature(root)
    details_by_feature = detail_slugs_by_feature(root)
    rows: list[dict[str, str]] = []
    for feature_id in sorted(features):
        feature = features[feature_id]
        title = (feature.get("title") or feature_id).strip()
        slug = slugify(f"{feature_id}-{title}")
        linked_detail_maps = details_by_feature.get(feature_id, [])
        candidate_slugs = [row.get("proposed_slug", "") for row in candidates_by_feature.get(feature_id, []) if row.get("proposed_slug")]
        artifacts = scenario_artifacts.get(feature_id, {})
        evidence_refs: list[str] = []
        for artifact_name in ("summary.md", "evidence.csv"):
            artifact = artifacts.get(artifact_name)
            if artifact:
                evidence_refs.append(artifact.relative_to(root).as_posix())
        feature_evidence = repo_path(root, f"analysis/features/{feature_id}/evidence.csv")
        if feature_evidence.exists():
            evidence_refs.append(feature_evidence.relative_to(root).as_posix())
        if linked_detail_maps:
            evidence_refs.append("analysis/detail-maps/index.csv")
        runtime = [path.relative_to(root).as_posix() for name, path in sorted(artifacts.items()) if name.startswith("runtime-")]
        runtime.extend(runtime_refs.get(feature_id, []))
        summary_text = first_sentence(feature.get("summary", "")) or f"Контур объединяет подтвержденные изменения {feature_id}: {title}."
        rows.append(
            {
                "contour_id": f"SCN-{len(rows) + 1:04d}",
                "slug": slug,
                "title": title,
                "contour_type": subject_type_for_feature(feature),
                "linked_features": feature_id,
                "primary_objects": feature.get("source_bucket", ""),
                "linked_detail_maps": join_refs(linked_detail_maps),
                "scenario_summary": summary_text,
                "migration_boundary": (
                    f"Переносить и проверять как единый контур `{title}`: основную модель, формы/команды, "
                    "права, регламентные операции, обмены и отчеты, которые доказательно связаны с этим сценарием."
                ),
                "why_this_is_one_contour": (
                    f"`{title}` является ревьюируемым контуром, потому что final feature {feature_id}, reverse-map "
                    "и detail maps связывают изменения в одну миграционную область."
                ),
                "why_not_technical_bucket": (
                    "Контур задан по смыслу миграции и пользовательскому/интеграционному назначению, "
                    "а технические группы объектов учитываются только как поддерживающие доказательства."
                ),
                "evidence_refs": join_refs(evidence_refs),
                "runtime_refs": join_refs(runtime),
                "technical_bucket_refs": join_refs(candidate_slugs),
                "status": "candidate",
                "confidence": feature.get("confidence", "") or "medium",
                "notes": "BF-контур является кандидатом на предметную классификацию; агент должен восстановить бизнес-сценарии до принятия готовой карточки.",
            }
        )
    write_csv_rows(subject_root(root) / "contours.csv", SUBJECT_CARD_CONTOURS_HEADER, rows)
    return {"contours": len(rows), "accepted": sum(1 for row in rows if row.get("status") == "accepted")}


def validate_subject_card_contours(root: Path) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    errors: list[str] = []
    path = subject_root(root) / "contours.csv"
    actual = path.read_text(encoding="utf-8-sig").splitlines()[:1] if path.exists() else []
    if actual != [SUBJECT_CARD_CONTOURS_HEADER]:
        errors.append("analysis/subject-cards/contours.csv: неверный заголовок CSV")
    rows = contour_rows(root)
    if not rows:
        errors.append("analysis/subject-cards/contours.csv: нет контуров")
    accepted = [row for row in rows if row.get("status") == "accepted"]
    required = (
        "contour_id",
        "slug",
        "title",
        "contour_type",
        "linked_features",
        "scenario_summary",
        "migration_boundary",
        "why_this_is_one_contour",
        "why_not_technical_bucket",
        "evidence_refs",
        "status",
        "confidence",
    )
    feature_ids = set(feature_by_id(root))
    covered_features: set[str] = set()
    for row in rows:
        label = row.get("contour_id") or row.get("slug") or "<empty>"
        if row.get("status") == "accepted":
            for field in required:
                if not str(row.get(field) or "").strip():
                    errors.append(f"{label}: не заполнено поле {field}")
            if row.get("contour_type") not in SUBJECT_CARD_TYPES:
                errors.append(f"{label}: недопустимый contour_type {row.get('contour_type')}")
            if generic_bucket_title(row.get("title", "")) and "не техническая корзина" not in row.get("why_not_technical_bucket", "").lower():
                errors.append(f"{label}: generic-название требует явного contour-обоснования")
            for field in ("scenario_summary", "migration_boundary", "why_this_is_one_contour", "why_not_technical_bucket"):
                if generic_template_text(row.get(field, "")):
                    errors.append(f"{label}: поле {field} похоже на шаблонный текст")
            for feature_id in split_refs(row.get("linked_features", "")):
                covered_features.add(feature_id)
            if not split_refs(row.get("technical_bucket_refs", "")):
                errors.append(f"{label}: нет technical_bucket_refs для связи с исходными кандидатами")
        else:
            for feature_id in split_refs(row.get("linked_features", "")):
                covered_features.add(feature_id)
            if row.get("status") not in SUBJECT_CARD_STATUSES:
                errors.append(f"{label}: недопустимый status {row.get('status')}")
    for feature_id in sorted(feature_ids - covered_features):
        errors.append(f"BF {feature_id}: нет contour-кандидата или accepted contour")
    return {"status": "ok" if not errors else "fail", "contours": len(rows), "accepted": len(accepted), "errors": errors}


def subject_type_for_detail_map(data: dict[str, Any]) -> str:
    map_type = str(data.get("type") or "").strip()
    return {
        "document": "business_document",
        "catalog": "reference_model",
        "route": "business_process",
        "scheduled_job": "background_automation",
        "rights": "access_model",
        "integration": "integration",
        "report": "business_process",
        "ui": "ui_surface",
    }.get(map_type, "business_process")


def subject_type_for_feature(row: dict[str, str]) -> str:
    classification = (row.get("classification") or "").strip().lower()
    title = (row.get("title") or "").strip().lower()
    if "access" in classification or "security" in classification or "права" in title or "полномоч" in title:
        return "access_model"
    if "integration" in classification or "синхрон" in title or "выгруз" in title:
        return "integration"
    if "automation" in classification or "регламент" in title or "фоновые" in title:
        return "background_automation"
    if "ui" in classification or "пользовательский ui" in title:
        return "ui_surface"
    if "reference data" in classification or "справоч" in title or "нормативно-справ" in title:
        return "reference_model"
    if "migration" in classification or "service" in classification or "миграцион" in title or "сервисн" in title:
        return "technical_support"
    return "business_process"


def primary_objects_from_detail_map(data: dict[str, Any]) -> list[str]:
    result: list[str] = []
    for section in ("attributes", "form_rules", "validations", "lifecycle", "rights", "scheduled_jobs", "ui", "integrations"):
        rows = data.get(section, [])
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                continue
            for key in ("object", "surface", "name", "system", "role"):
                value = str(row.get(key) or "").strip()
                if value:
                    result.append(value)
                    break
            if len(result) >= 12:
                return split_refs(join_refs(result))
    return split_refs(join_refs(result))


def detail_slugs_by_feature(root: Path) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for _, data in read_detail_maps(root):
        slug = str(data.get("slug") or "").strip()
        if not slug:
            continue
        for feature_id in split_refs(data.get("linked_features", [])):
            result.setdefault(feature_id, []).append(slug)
    return {key: split_refs(join_refs(value)) for key, value in result.items()}


def card_reason(slug: str, subject_type: str, origin_layer: str) -> str:
    if origin_layer == "detail_map":
        return "Detail-map описывает самостоятельный предметный объект с собственными реквизитами, поведением и рисками перехода."
    if subject_type == "business_document":
        return "Документ или вид документа образует самостоятельную предметную доработку с собственными реквизитами, поведением и рисками перехода."
    if subject_type == "business_process":
        return "Бизнес-процесс образует самостоятельную предметную доработку, если его сценарии, формы, состояния и фоновые операции ревьюятся вместе."
    if subject_type == "technical_support":
        return "BF выглядит как поддерживающий технический слой; нужна отдельная классификация перед созданием карточки."
    return "BF пока является контейнером фактов; нужна evidence-based классификация split/merge перед созданием предметной карточки."


def coverage_scope_for_status(status: str, card_path: str, decision: str = "") -> str:
    if card_path:
        return "covered"
    if decision == "supporting" or status == "supporting":
        return "supporting"
    if decision == "merge" or status == "merged_into_other":
        return "shared"
    if decision == "reject" or status == "rejected":
        return "rejected"
    if status == "candidate" or decision == "split":
        return "unclassified"
    return "unclassified"


def card_status(payload: dict[str, Any], evidence_rows: list[dict[str, str]], gaps: list[dict[str, str]]) -> str:
    if any(row.get("blocking") == "true" and row.get("status") != "closed" for row in gaps):
        for row in gaps:
            if row.get("blocking") == "true" and row.get("status") != "closed":
                if row.get("needed_source") == "runtime":
                    return "needs_runtime_data"
                if row.get("needed_source") == "ui":
                    return "needs_ui_check"
        return "needs_static_analysis"
    required = ("summary", "identification", "key_conclusion", "upgrade_risk")
    if all(str(payload.get(field) or "").strip() for field in required) and evidence_rows:
        return "ready_for_review"
    return "draft"


def row_claim(section: str, row: dict[str, str]) -> str:
    for key in ("relation", "description", "validation", "action", "rule", "claim", "question", "evidence", "summary"):
        value = row.get(key, "")
        if value:
            return value
    return f"Строка секции {section}"


def evidence_from_sections(sections: dict[str, list[dict[str, str]]], fallback_source: str, linked_features: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for section, section_rows in sections.items():
        if section == "open_questions":
            continue
        for row in section_rows:
            source = row.get("source") or row.get("source_path") or fallback_source
            line = row.get("line") or row.get("line_start") or ""
            rows.append(
                {
                    "evidence_id": f"E-{len(rows) + 1:04d}",
                    "section": section,
                    "claim": row_claim(section, row),
                    "source_type": "artifact" if source == fallback_source else "source",
                    "source_path": source,
                    "line": line,
                    "linked_diff_id": row.get("id", ""),
                    "linked_feature_id": ";".join(linked_features),
                    "confidence": row.get("confidence", ""),
                    "notes": "",
                }
            )
    if not rows and fallback_source:
        rows.append(
            {
                "evidence_id": "E-0001",
                "section": "sources",
                "claim": "Карточка создана из существующего артефакта анализа.",
                "source_type": "artifact",
                "source_path": fallback_source,
                "line": "",
                "linked_diff_id": "",
                "linked_feature_id": ";".join(linked_features),
                "confidence": "",
                "notes": "",
            }
        )
    return rows


def gaps_for_payload(payload: dict[str, Any], evidence_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    gaps: list[dict[str, str]] = []
    for field, question in (
        ("summary", "Нужно заполнить сводку предметной доработки."),
        ("identification", "Нужно описать, как предметная доработка идентифицируется в исходниках или ИБ."),
        ("key_conclusion", "Нужно сформулировать ключевой вывод на основании доказательств."),
        ("upgrade_risk", "Нужно описать риск перехода на целевой релиз."),
    ):
        if not str(payload.get(field) or "").strip():
            gaps.append(
                {
                    "gap_id": f"G-{len(gaps) + 1:04d}",
                    "section": field,
                    "question": question,
                    "needed_source": "static",
                    "status": "open",
                    "blocking": "true",
                    "notes": "",
                }
            )
    if not evidence_rows:
        gaps.append(
            {
                "gap_id": f"G-{len(gaps) + 1:04d}",
                "section": "sources",
                "question": "Нужно добавить хотя бы одно доказательство карточки.",
                "needed_source": "static",
                "status": "open",
                "blocking": "true",
                "notes": "",
            }
        )
    return gaps


def subject_from_detail_map(root: Path, candidate: dict[str, str], path: Path, data: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, str]], list[dict[str, str]]]:
    slug = candidate["proposed_slug"]
    sections = empty_sections()
    for section in SUBJECT_CARD_SECTIONS:
        section_rows = normalize_section_rows(data.get(section, []))
        if section != "open_questions":
            for row in section_rows:
                row.setdefault("claim", row_claim(section, row))
        sections[section] = section_rows
    linked_features = split_refs(data.get("linked_features", []))
    source_path = path.relative_to(root).as_posix()
    payload: dict[str, Any] = {
        "schema_version": "subject-card/v1",
        "slug": slug,
        "title": str(data.get("title") or candidate.get("title") or slug),
        "status": "draft",
        "confidence": str(data.get("confidence") or candidate.get("confidence") or "medium"),
        "subject_type": candidate.get("subject_type") or subject_type_for_detail_map(data),
        "origin_layer": "detail_map",
        "primary_objects": split_refs(candidate.get("primary_objects", "")) or primary_objects_from_detail_map(data),
        "coverage_scope": "covered",
        "why_separate_card": card_reason(slug, candidate.get("subject_type") or subject_type_for_detail_map(data), "detail_map"),
        "merge_into": "",
        "split_from": "",
        "identification": str(data.get("identification") or ""),
        "summary": str(data.get("summary") or ""),
        "key_conclusion": str(data.get("key_conclusion") or ""),
        "upgrade_risk": str(data.get("upgrade_risk") or ""),
        "runtime_data_needed": str(data.get("runtime_data_needed") or ""),
        "review_status": str(data.get("review_status") or ""),
        "linked_features": linked_features,
        "linked_detail_maps": [str(data.get("slug") or path.parent.name)],
        "source_mode": "detail_map_artifact",
        "source_artifacts": [source_path],
        "generated_at": utc_now_iso(),
        "sections": sections,
    }
    evidence_rows = evidence_from_sections(sections, source_path, linked_features)
    gaps = gaps_for_payload(payload, evidence_rows)
    payload["status"] = card_status(payload, evidence_rows, gaps)
    return payload, evidence_rows, gaps


def subject_from_feature(root: Path, candidate: dict[str, str], row: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]], list[dict[str, str]]]:
    slug = candidate["proposed_slug"]
    feature_id = (row.get("feature_id") or "").strip()
    sections = empty_sections()
    evidence_path = repo_path(root, row.get("evidence_pack_path", "") or f"analysis/features/{feature_id}") / "evidence.csv"
    feature_evidence = non_empty_rows(read_csv_rows(evidence_path))
    source_artifacts = [evidence_path.relative_to(root).as_posix()] if evidence_path.exists() else []
    evidence_rows: list[dict[str, str]] = []
    for row_index, evidence in enumerate(feature_evidence[:20], start=1):
        evidence_rows.append(
            {
                "evidence_id": f"E-{row_index:04d}",
                "section": "sources",
                "claim": evidence.get("summary", ""),
                "source_type": evidence.get("source_kind", ""),
                "source_path": evidence.get("source_path", ""),
                "line": evidence.get("line_start", ""),
                "linked_diff_id": "",
                "linked_feature_id": feature_id,
                "confidence": evidence.get("confidence", ""),
                "notes": evidence.get("notes", ""),
            }
        )
    title = candidate.get("title") or row.get("title") or slug
    summary = row.get("summary", "")
    payload: dict[str, Any] = {
        "schema_version": "subject-card/v1",
        "slug": slug,
        "title": title,
        "status": "draft",
        "confidence": row.get("confidence") or candidate.get("confidence") or "medium",
        "subject_type": candidate.get("subject_type") or subject_type_for_feature(row),
        "origin_layer": "BF",
        "primary_objects": split_refs(candidate.get("primary_objects", "")),
        "coverage_scope": "covered" if candidate.get("proposed_action") == "accept" else "unclassified",
        "why_separate_card": card_reason(slug, candidate.get("subject_type") or subject_type_for_feature(row), "BF"),
        "merge_into": "",
        "split_from": "",
        "identification": f"Карточка выделена из BF-контейнера {feature_id}; источник фактов: финальная карта фич и evidence pack.",
        "summary": summary,
        "key_conclusion": (
            f"Доработка `{title}` подтверждена статическим clean diff и сгруппирована в {feature_id}. "
            "Карточка готова к первичному аналитическому ревью; границы можно уточнять через subject-card registry."
        ),
        "upgrade_risk": (
            "Перед переходом на целевой релиз нужно сопоставить связанные объекты и сценарии с новой типовой конфигурацией; "
            "если типовой механизм закрывает сценарий, доработку можно пометить как merge/supporting в реестре subject-card."
        ),
        "runtime_data_needed": "",
        "review_status": "Автокарточка из BF: готова к первичному ревью аналитиком, требует ручного уточнения границ при необходимости.",
        "linked_features": [feature_id] if feature_id else [],
        "linked_detail_maps": [],
        "source_mode": "feature_candidate",
        "source_artifacts": source_artifacts,
        "generated_at": utc_now_iso(),
        "sections": sections,
    }
    gaps = gaps_for_payload(payload, evidence_rows)
    payload["status"] = card_status(payload, evidence_rows, gaps)
    return payload, evidence_rows, gaps


def subject_from_contour(root: Path, contour: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]], list[dict[str, str]]]:
    slug = contour["slug"]
    linked_features = split_refs(contour.get("linked_features", ""))
    linked_detail_maps = split_refs(contour.get("linked_detail_maps", ""))
    evidence_refs = split_refs(contour.get("evidence_refs", ""))
    runtime_refs = split_refs(contour.get("runtime_refs", ""))
    sections = empty_sections()
    sections["sources"] = []
    for ref in evidence_refs[:12]:
        sections["sources"].append(
            {
                "claim": f"Контур подтвержден артефактом `{ref}`.",
                "source": ref,
                "line": "",
                "confidence": contour.get("confidence", ""),
            }
        )
    for ref in runtime_refs[:8]:
        sections["lifecycle"].append(
            {
                "claim": f"Runtime/ИБ-проверка учтена в границах контура: `{ref}`.",
                "source": ref,
                "line": "",
                "confidence": contour.get("confidence", ""),
            }
        )
    payload: dict[str, Any] = {
        "schema_version": "subject-card/v1",
        "slug": slug,
        "title": contour.get("title", "") or slug,
        "status": "draft",
        "confidence": contour.get("confidence", "") or "medium",
        "subject_type": contour.get("contour_type", "") or "business_process",
        "origin_layer": "contour",
        "primary_objects": split_refs(contour.get("primary_objects", "")),
        "coverage_scope": "covered",
        "why_separate_card": contour.get("why_this_is_one_contour", ""),
        "why_not_technical_bucket": contour.get("why_not_technical_bucket", ""),
        "migration_boundary": contour.get("migration_boundary", ""),
        "merge_into": "",
        "split_from": "",
        "identification": contour.get("scenario_summary", ""),
        "summary": contour.get("scenario_summary", ""),
        "key_conclusion": contour.get("why_this_is_one_contour", ""),
        "upgrade_risk": contour.get("migration_boundary", ""),
        "runtime_data_needed": "Runtime/ИБ-проверки закрыты и учтены как доказательства." if runtime_refs else "",
        "review_status": "Контурная карточка готова к аналитическому ревью; технические detail maps являются подложкой.",
        "linked_features": linked_features,
        "linked_detail_maps": linked_detail_maps,
        "source_mode": "contour_map",
        "source_artifacts": split_refs(join_refs(evidence_refs + runtime_refs + ["analysis/subject-cards/contours.csv"])),
        "generated_at": utc_now_iso(),
        "sections": sections,
    }
    evidence_rows = evidence_from_sections(sections, "analysis/subject-cards/contours.csv", linked_features)
    gaps = gaps_for_payload(payload, evidence_rows)
    payload["status"] = card_status(payload, evidence_rows, gaps)
    return payload, evidence_rows, gaps


def write_review(path: Path, payload: dict[str, Any], evidence_rows: list[dict[str, str]], gaps: list[dict[str, str]]) -> None:
    open_gaps = [row for row in gaps if row.get("status") != "closed"]
    lines = [
        f"# Ревью предметной карточки: {payload.get('title', payload.get('slug', ''))}",
        "",
        f"- Статус: `{payload.get('status', '')}`",
        f"- Уверенность: `{payload.get('confidence', '')}`",
        f"- Доказательств: {len(evidence_rows)}",
        f"- Открытых gaps: {len(open_gaps)}",
        "",
        "## Ключевой вывод",
        "",
        str(payload.get("key_conclusion") or "Ключевой вывод не заполнен."),
        "",
        "## Граница переноса",
        "",
        str(payload.get("migration_boundary") or payload.get("upgrade_risk") or "Граница переноса не заполнена."),
        "",
        "## Следующий шаг",
        "",
    ]
    if open_gaps:
        lines.append("Закрыть gaps из `gaps.csv`, затем повторить `subject-card refine` и `subject-card validate`.")
    else:
        lines.append("Карточка готова к аналитическому ревью.")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


def write_card_bundle(root: Path, payload: dict[str, Any], evidence_rows: list[dict[str, str]], gaps: list[dict[str, str]]) -> None:
    card_dir = subject_root(root) / "cards" / payload["slug"]
    card_dir.mkdir(parents=True, exist_ok=True)
    (card_dir / "subject-card.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    write_csv_rows(card_dir / "evidence.csv", SUBJECT_CARD_EVIDENCE_HEADER, evidence_rows)
    write_csv_rows(card_dir / "gaps.csv", SUBJECT_CARD_GAPS_HEADER, gaps)
    write_review(card_dir / "review.md", payload, evidence_rows, gaps)


def candidate_rows(root: Path) -> list[dict[str, str]]:
    path = subject_root(root) / "candidates.csv"
    rows = non_empty_rows(read_csv_rows(path))
    if rows:
        return rows
    discover_subject_cards(root)
    return non_empty_rows(read_csv_rows(path))


def card_payloads(root: Path) -> dict[str, tuple[Path, dict[str, Any], int, int]]:
    result: dict[str, tuple[Path, dict[str, Any], int, int]] = {}
    for card_dir in card_dirs(root):
        try:
            payload = json.loads((card_dir / "subject-card.json").read_text(encoding="utf-8"))
        except Exception:
            continue
        slug = str(payload.get("slug") or card_dir.name).strip()
        if not slug:
            continue
        evidence_count = len(non_empty_rows(read_csv_rows(card_dir / "evidence.csv")))
        gap_count = len([row for row in non_empty_rows(read_csv_rows(card_dir / "gaps.csv")) if row.get("status") != "closed"])
        result[slug] = (card_dir, payload, evidence_count, gap_count)
    return result


def classification_rows(root: Path) -> list[dict[str, str]]:
    path = subject_root(root) / "classification.csv"
    rows = non_empty_rows(read_csv_rows(path))
    if rows:
        return rows
    classify_subject_cards(root)
    return non_empty_rows(read_csv_rows(path))


def classify_subject_cards(root: Path) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    contours = contour_rows(root)
    if contours:
        candidates = candidate_rows(root)
        candidates_by_slug = {row.get("proposed_slug", ""): row for row in candidates}
        accepted = [row for row in contours if row.get("status") == "accepted"]
        covered_technical: set[str] = set()
        rows: list[dict[str, str]] = []
        for contour in accepted:
            slug = contour.get("slug", "")
            rows.append(
                {
                    "candidate_id": contour.get("contour_id", ""),
                    "proposed_slug": slug,
                    "decision": "accept",
                    "subject_type": contour.get("contour_type", "") or "business_process",
                    "registry_slug": slug,
                    "merge_into": "",
                    "split_from": "",
                    "why_separate_card": contour.get("why_this_is_one_contour", ""),
                    "status": "ready_for_review",
                    "confidence": contour.get("confidence", "") or "medium",
                    "notes": "Принято из contour-map; технические корзины связаны через technical_bucket_refs.",
                }
            )
            for technical_slug in split_refs(contour.get("technical_bucket_refs", "")):
                if technical_slug:
                    covered_technical.add(technical_slug)
        for candidate in candidates:
            slug = candidate.get("proposed_slug", "")
            if not slug or slug in {row.get("registry_slug") for row in rows}:
                continue
            if slug in covered_technical:
                decision = "supporting"
                status = "supporting"
                merge_into = ""
                notes = "Техническая корзина покрыта accepted contour и не является финальной карточкой."
            else:
                decision = "split"
                status = "candidate"
                merge_into = ""
                notes = "Кандидат не покрыт accepted contour; требуется контурное решение."
            rows.append(
                {
                    "candidate_id": candidate.get("candidate_id", ""),
                    "proposed_slug": slug,
                    "decision": decision,
                    "subject_type": candidate.get("subject_type", "") or "business_process",
                    "registry_slug": slug,
                    "merge_into": merge_into,
                    "split_from": "",
                    "why_separate_card": candidate.get("discovery_basis", ""),
                    "status": status,
                    "confidence": candidate.get("confidence", "") or "medium",
                    "notes": notes,
                }
            )
        write_csv_rows(subject_root(root) / "classification.csv", SUBJECT_CARD_CLASSIFICATION_HEADER, rows)
        return {"classified": len(rows), "accepted": sum(1 for row in rows if row.get("decision") == "accept")}

    candidates = candidate_rows(root)
    cards = card_payloads(root)
    rows: list[dict[str, str]] = []
    for candidate in candidates:
        slug = (candidate.get("proposed_slug") or "").strip()
        if not slug:
            continue
        candidate_action = (candidate.get("proposed_action") or "split").strip()
        if candidate_action not in SUBJECT_CARD_ACTIONS:
            candidate_action = "split"
        card_payload = cards.get(slug, (Path(), {}, 0, 0))[1]
        has_card = bool(card_payload)
        decision = "accept" if has_card or candidate_action == "accept" else candidate_action
        subject_type = (card_payload.get("subject_type") or candidate.get("subject_type") or "business_process").strip()
        if subject_type not in SUBJECT_CARD_TYPES:
            subject_type = "business_process"
        status = str(card_payload.get("status") or candidate.get("status") or "candidate").strip()
        if not has_card and decision in {"split", "merge", "supporting"}:
            status = "candidate" if decision == "split" else decision
        rows.append(
            {
                "candidate_id": candidate.get("candidate_id", ""),
                "proposed_slug": slug,
                "decision": decision,
                "subject_type": subject_type,
                "registry_slug": slug,
                "merge_into": str(card_payload.get("merge_into") or ""),
                "split_from": str(card_payload.get("split_from") or ""),
                "why_separate_card": str(card_payload.get("why_separate_card") or card_reason(slug, subject_type, "detail_map" if candidate.get("linked_detail_maps") and candidate.get("source", "").endswith("detail-map.json") else "BF")),
                "status": status,
                "confidence": str(card_payload.get("confidence") or candidate.get("confidence") or "medium"),
                "notes": candidate.get("notes", ""),
            }
        )
    write_csv_rows(subject_root(root) / "classification.csv", SUBJECT_CARD_CLASSIFICATION_HEADER, rows)
    return {"classified": len(rows), "accepted": sum(1 for row in rows if row.get("decision") == "accept")}


def hydrate_card_payload(payload: dict[str, Any], registry_row: dict[str, str]) -> bool:
    changed = False
    updates: dict[str, Any] = {
        "subject_type": registry_row.get("subject_type", ""),
        "origin_layer": registry_row.get("origin_layer", ""),
        "primary_objects": split_refs(registry_row.get("primary_objects", "")),
        "coverage_scope": registry_row.get("coverage_scope", ""),
        "why_separate_card": registry_row.get("why_separate_card", ""),
        "merge_into": registry_row.get("merge_into", ""),
        "split_from": registry_row.get("split_from", ""),
    }
    for key, value in updates.items():
        if value in ("", []):
            continue
        if payload.get(key) != value:
            payload[key] = value
            changed = True
    return changed


def registry_row_from_card(
    root: Path,
    slug: str,
    card_dir: Path,
    payload: dict[str, Any],
    evidence_count: int,
    gap_count: int,
    classification: dict[str, str],
) -> dict[str, str]:
    linked_features = split_refs(payload.get("linked_features", [])) or split_refs(classification.get("linked_features", ""))
    linked_detail_maps = split_refs(payload.get("linked_detail_maps", []))
    subject_type = str(payload.get("subject_type") or classification.get("subject_type") or "business_process").strip()
    if subject_type not in SUBJECT_CARD_TYPES:
        subject_type = "business_process"
    source_mode = str(payload.get("source_mode") or "")
    if source_mode.startswith("feature_candidate"):
        origin_layer = "BF"
    elif source_mode == "detail_map_artifact":
        origin_layer = "detail_map"
    else:
        origin_layer = str(payload.get("origin_layer") or ("detail_map" if linked_detail_maps else "BF"))
    primary_objects = split_refs(payload.get("primary_objects", []))
    if primary_objects == ["clean_rebase_diff"] and linked_features:
        feature = feature_by_id(root).get(linked_features[0], {})
        source_bucket = str(feature.get("source_bucket") or "").strip()
        if source_bucket:
            primary_objects = [source_bucket]
    why = str(payload.get("why_separate_card") or classification.get("why_separate_card") or card_reason(slug, subject_type, origin_layer))
    return {
        "slug": slug,
        "title": str(payload.get("title") or slug),
        "subject_type": subject_type,
        "status": str(payload.get("status") or "draft"),
        "confidence": str(payload.get("confidence") or classification.get("confidence") or ""),
        "origin_layer": origin_layer,
        "owner_feature": linked_features[0] if linked_features else "",
        "linked_features": join_refs(linked_features),
        "linked_detail_maps": join_refs(linked_detail_maps),
        "primary_objects": join_refs(primary_objects),
        "coverage_scope": str(payload.get("coverage_scope") or "covered"),
        "why_separate_card": why,
        "merge_into": str(payload.get("merge_into") or classification.get("merge_into") or ""),
        "split_from": str(payload.get("split_from") or classification.get("split_from") or ""),
        "card_path": (card_dir / "subject-card.json").relative_to(root).as_posix(),
        "evidence_count": str(evidence_count),
        "gap_count": str(gap_count),
        "review_notes": str(payload.get("review_status") or classification.get("notes") or ""),
    }


def registry_row_from_candidate(candidate: dict[str, str], classification: dict[str, str]) -> dict[str, str]:
    slug = classification.get("registry_slug") or candidate.get("proposed_slug", "")
    subject_type = classification.get("subject_type") or candidate.get("subject_type") or "business_process"
    if subject_type not in SUBJECT_CARD_TYPES:
        subject_type = "business_process"
    decision = classification.get("decision") or candidate.get("proposed_action") or "split"
    status = classification.get("status") or candidate.get("status") or "candidate"
    origin_layer = "detail_map" if candidate.get("source", "").endswith("detail-map.json") else "BF"
    coverage_scope = coverage_scope_for_status(status, "", decision)
    return {
        "slug": slug,
        "title": candidate.get("title", "") or slug,
        "subject_type": subject_type,
        "status": status,
        "confidence": classification.get("confidence") or candidate.get("confidence") or "medium",
        "origin_layer": origin_layer,
        "owner_feature": split_refs(candidate.get("linked_features", ""))[0] if split_refs(candidate.get("linked_features", "")) else "",
        "linked_features": candidate.get("linked_features", ""),
        "linked_detail_maps": candidate.get("linked_detail_maps", ""),
        "primary_objects": candidate.get("primary_objects", ""),
        "coverage_scope": coverage_scope,
        "why_separate_card": classification.get("why_separate_card") or card_reason(slug, subject_type, origin_layer),
        "merge_into": classification.get("merge_into", ""),
        "split_from": classification.get("split_from", ""),
        "card_path": "",
        "evidence_count": "0",
        "gap_count": "0",
        "review_notes": classification.get("notes") or candidate.get("notes", ""),
    }


def registry_row_from_classification(classification: dict[str, str]) -> dict[str, str]:
    slug = classification.get("registry_slug") or classification.get("proposed_slug", "")
    decision = classification.get("decision", "")
    status = classification.get("status", "") or ("ready_for_review" if decision == "accept" else "candidate")
    return {
        "slug": slug,
        "title": slug,
        "subject_type": classification.get("subject_type", "") or "business_process",
        "status": status,
        "confidence": classification.get("confidence", "") or "medium",
        "origin_layer": "contour" if decision == "accept" else "classification",
        "owner_feature": "",
        "linked_features": "",
        "linked_detail_maps": "",
        "primary_objects": "",
        "coverage_scope": coverage_scope_for_status(status, "", decision),
        "why_separate_card": classification.get("why_separate_card", ""),
        "merge_into": classification.get("merge_into", ""),
        "split_from": classification.get("split_from", ""),
        "card_path": "",
        "evidence_count": "0",
        "gap_count": "0",
        "review_notes": classification.get("notes", ""),
    }


def registry_row_from_contour(contour: dict[str, str], classification: dict[str, str]) -> dict[str, str]:
    slug = contour.get("slug", "")
    return {
        "slug": slug,
        "title": contour.get("title", "") or slug,
        "subject_type": contour.get("contour_type", "") or classification.get("subject_type", "") or "business_process",
        "status": classification.get("status", "") or "ready_for_review",
        "confidence": contour.get("confidence", "") or classification.get("confidence", "") or "medium",
        "origin_layer": "contour",
        "owner_feature": split_refs(contour.get("linked_features", ""))[0] if split_refs(contour.get("linked_features", "")) else "",
        "linked_features": contour.get("linked_features", ""),
        "linked_detail_maps": contour.get("linked_detail_maps", ""),
        "primary_objects": contour.get("primary_objects", ""),
        "coverage_scope": "covered",
        "why_separate_card": contour.get("why_this_is_one_contour", "") or classification.get("why_separate_card", ""),
        "merge_into": classification.get("merge_into", ""),
        "split_from": classification.get("split_from", ""),
        "card_path": "",
        "evidence_count": "0",
        "gap_count": "0",
        "review_notes": classification.get("notes", "") or contour.get("notes", ""),
    }


def coverage_relation(row: dict[str, str]) -> str:
    if row.get("card_path"):
        return "covered_by_subject_card"
    status = row.get("status", "")
    scope = row.get("coverage_scope", "")
    if status == "supporting" or scope == "supporting":
        return "supporting"
    if status == "merged_into_other" or scope == "shared":
        return "shared"
    if status == "rejected" or scope == "rejected":
        return "rejected"
    if status == "candidate":
        return "unclassified"
    return scope or "unclassified"


def build_subject_card_coverage(root: Path, registry_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    features = feature_by_id(root)
    for feature_id, feature in features.items():
        linked = [row for row in registry_rows if feature_id in split_refs(row.get("linked_features", ""))]
        if not linked:
            rows.append(
                {
                    "source_kind": "BF",
                    "source_id": feature_id,
                    "feature_id": feature_id,
                    "detail_map_slug": "",
                    "subject_card_slug": "",
                    "relation": "unclassified",
                    "confidence": feature.get("confidence", ""),
                    "notes": "BF пока не связан с subject card или классификационным решением.",
                }
            )
            continue
        for item in linked:
            rows.append(
                {
                    "source_kind": "BF",
                    "source_id": feature_id,
                    "feature_id": feature_id,
                    "detail_map_slug": "",
                    "subject_card_slug": item.get("slug", ""),
                    "relation": coverage_relation(item),
                    "confidence": item.get("confidence", ""),
                    "notes": item.get("coverage_scope", ""),
                }
            )

    for _, data in read_detail_maps(root):
        slug = str(data.get("slug") or "").strip()
        if not slug:
            continue
        linked = [row for row in registry_rows if slug in split_refs(row.get("linked_detail_maps", ""))]
        linked_features = split_refs(data.get("linked_features", []))
        if not linked:
            rows.append(
                {
                    "source_kind": "detail_map",
                    "source_id": slug,
                    "feature_id": linked_features[0] if linked_features else "",
                    "detail_map_slug": slug,
                    "subject_card_slug": "",
                    "relation": "technical_support" if str(data.get("generation_mode") or "") == "generated" else "unclassified",
                    "confidence": str(data.get("confidence") or ""),
                    "notes": "Generated detail-map является технической подложкой." if str(data.get("generation_mode") or "") == "generated" else "Detail-map пока не связан с subject card.",
                }
            )
            continue
        for item in linked:
            rows.append(
                {
                    "source_kind": "detail_map",
                    "source_id": slug,
                    "feature_id": linked_features[0] if linked_features else item.get("owner_feature", ""),
                    "detail_map_slug": slug,
                    "subject_card_slug": item.get("slug", ""),
                    "relation": coverage_relation(item),
                    "confidence": item.get("confidence", ""),
                    "notes": item.get("coverage_scope", ""),
                }
            )
    write_csv_rows(subject_root(root) / "coverage.csv", SUBJECT_CARD_COVERAGE_HEADER, rows)
    return rows


def build_subject_card_registry(root: Path) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    candidates = {row.get("proposed_slug", ""): row for row in candidate_rows(root)}
    classifications = {row.get("registry_slug") or row.get("proposed_slug", ""): row for row in classification_rows(root)}
    contours = accepted_contours(root)
    cards = card_payloads(root)
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for slug, (card_dir, payload, evidence_count, gap_count) in sorted(cards.items()):
        row = registry_row_from_card(root, slug, card_dir, payload, evidence_count, gap_count, classifications.get(slug, {}))
        if hydrate_card_payload(payload, row):
            (card_dir / "subject-card.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
        rows.append(row)
        seen.add(slug)
    for slug, candidate in sorted(candidates.items()):
        if not slug or slug in seen:
            continue
        rows.append(registry_row_from_candidate(candidate, classifications.get(slug, {})))
        seen.add(slug)
    for slug, classification in sorted(classifications.items()):
        if not slug or slug in seen:
            continue
        if slug in contours:
            rows.append(registry_row_from_contour(contours[slug], classification))
        else:
            rows.append(registry_row_from_classification(classification))
        seen.add(slug)
    write_csv_rows(subject_root(root) / "registry.csv", SUBJECT_CARD_REGISTRY_HEADER, rows)
    coverage = build_subject_card_coverage(root, rows)
    return {"registry_rows": len(rows), "coverage_rows": len(coverage), "cards": len(cards)}


def refresh_registry(root: Path) -> None:
    build_subject_card_registry(root)


def seed_subject_cards(root: Path, all_candidates: bool = False, card: str = "", from_registry: bool = False) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    if from_registry:
        registry_rows = non_empty_rows(read_csv_rows(subject_root(root) / "registry.csv"))
        approved_registry = {
            row.get("slug", ""): row
            for row in registry_rows
            if (row.get("card_path") or "").strip() == "" and row.get("status") in {"accepted", "ready_for_review", "draft"}
        }
        contour_map = accepted_contours(root)
        candidates = []
        seen_registry: set[str] = set()
        for row in candidate_rows(root):
            slug = row.get("proposed_slug", "")
            registry_row = approved_registry.get(slug)
            if not registry_row:
                continue
            seen_registry.add(slug)
            candidate = dict(row)
            candidate["status"] = "accepted"
            candidate["proposed_action"] = "accept"
            for key in ("title", "subject_type", "confidence", "linked_features", "linked_detail_maps", "primary_objects"):
                if registry_row.get(key):
                    candidate[key] = registry_row[key]
            candidates.append(candidate)
        for slug, registry_row in sorted(approved_registry.items()):
            if slug in seen_registry and slug not in contour_map:
                continue
            if slug in contour_map:
                contour = contour_map[slug]
                candidates.append(
                    {
                        "candidate_id": contour.get("contour_id", ""),
                        "title": contour.get("title", ""),
                        "proposed_slug": slug,
                        "source": "analysis/subject-cards/contours.csv",
                        "discovery_basis": contour.get("why_this_is_one_contour", ""),
                        "linked_features": contour.get("linked_features", ""),
                        "linked_detail_maps": contour.get("linked_detail_maps", ""),
                        "primary_objects": contour.get("primary_objects", ""),
                        "subject_type": contour.get("contour_type", "") or registry_row.get("subject_type", ""),
                        "confidence": contour.get("confidence", "") or registry_row.get("confidence", ""),
                        "proposed_action": "accept",
                        "status": "accepted",
                        "notes": "Контурная карточка из contours.csv.",
                    }
                )
    else:
        candidates = candidate_rows(root)
    detail_maps = detail_map_by_slug(root)
    features = feature_by_id(root)
    contours = accepted_contours(root)
    written: list[str] = []
    skipped: list[str] = []
    for candidate in candidates:
        slug = candidate.get("proposed_slug", "").strip()
        if card and slug != card:
            continue
        status = candidate.get("status", "").strip()
        if status != "accepted" and not all_candidates:
            skipped.append(slug)
            continue
        source = candidate.get("source", "")
        payload: dict[str, Any]
        evidence_rows: list[dict[str, str]]
        gaps: list[dict[str, str]]
        if slug in detail_maps:
            path, data = detail_maps[slug]
            payload, evidence_rows, gaps = subject_from_detail_map(root, candidate, path, data)
        elif slug in contours or source == "analysis/subject-cards/contours.csv":
            contour = contours.get(slug, {})
            if not contour:
                skipped.append(slug)
                continue
            payload, evidence_rows, gaps = subject_from_contour(root, contour)
        elif source.endswith("final-feature-map.csv") or source.endswith("feature-map.csv"):
            feature_id = split_refs(candidate.get("linked_features", ""))[0] if split_refs(candidate.get("linked_features", "")) else ""
            feature = features.get(feature_id, {})
            payload, evidence_rows, gaps = subject_from_feature(root, candidate, feature)
        else:
            skipped.append(slug)
            continue
        write_card_bundle(root, payload, evidence_rows, gaps)
        written.append(slug)
    refresh_registry(root)
    return {"written": written, "skipped": skipped}


def load_subject_card(root: Path, slug: str) -> tuple[dict[str, Any], list[dict[str, str]], list[dict[str, str]]]:
    card_dir = subject_root(root) / "cards" / slug
    payload = json.loads((card_dir / "subject-card.json").read_text(encoding="utf-8"))
    evidence_rows = non_empty_rows(read_csv_rows(card_dir / "evidence.csv"))
    gaps = non_empty_rows(read_csv_rows(card_dir / "gaps.csv"))
    return payload, evidence_rows, gaps


def refine_subject_card(root: Path, card: str) -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    card_path = subject_root(root) / "cards" / card / "subject-card.json"
    if not card_path.exists():
        seed_subject_cards(root, card=card)
    payload, evidence_rows, gaps = load_subject_card(root, card)
    gaps = gaps_for_payload(payload, evidence_rows)
    payload["status"] = card_status(payload, evidence_rows, gaps)
    payload["generated_at"] = utc_now_iso()
    write_card_bundle(root, payload, evidence_rows, gaps)
    refresh_registry(root)
    return {"card": card, "status": payload["status"], "evidence_count": len(evidence_rows), "gap_count": len(gaps)}


def card_dirs(root: Path) -> list[Path]:
    cards_root = subject_root(root) / "cards"
    if not cards_root.exists():
        return []
    return sorted(path for path in cards_root.iterdir() if (path / "subject-card.json").exists())


def validate_subject_cards(root: Path, card: str = "") -> dict[str, Any]:
    root = root.resolve()
    ensure_scaffold(root)
    errors: list[str] = []
    for relative, header in (
        ("candidates.csv", SUBJECT_CARD_CANDIDATES_HEADER),
        ("contours.csv", SUBJECT_CARD_CONTOURS_HEADER),
        ("classification.csv", SUBJECT_CARD_CLASSIFICATION_HEADER),
        ("registry.csv", SUBJECT_CARD_REGISTRY_HEADER),
        ("coverage.csv", SUBJECT_CARD_COVERAGE_HEADER),
    ):
        path = subject_root(root) / relative
        if not path.exists():
            errors.append(f"Нет артефакта subject-card: {path.relative_to(root).as_posix()}")
            continue
        actual = path.read_text(encoding="utf-8-sig").splitlines()[:1]
        if actual != [header]:
            errors.append(f"{path.relative_to(root).as_posix()}: неверный заголовок CSV")
    registry_rows = non_empty_rows(read_csv_rows(subject_root(root) / "registry.csv"))
    coverage_rows = non_empty_rows(read_csv_rows(subject_root(root) / "coverage.csv"))
    contours = contour_rows(root)
    accepted_contour_by_slug = {row.get("slug", ""): row for row in contours if row.get("status") == "accepted"}
    contour_validation = validate_subject_card_contours(root)
    errors.extend(contour_validation["errors"])
    classifications = non_empty_rows(read_csv_rows(subject_root(root) / "classification.csv"))
    accepted_classifications = [row for row in classifications if row.get("decision") == "accept"]
    if classifications and len(accepted_classifications) == len(classifications) and len(classifications) > len(accepted_contour_by_slug):
        errors.append("classification.csv: все кандидаты приняты как accept; нужен contour-layer с supporting/merge/reject для технических корзин")
    registry_by_slug = {row.get("slug", ""): row for row in registry_rows}
    for row in registry_rows:
        subject_type = row.get("subject_type", "")
        if subject_type not in SUBJECT_CARD_TYPES:
            errors.append(f"registry.csv: недопустимый subject_type для {row.get('slug', '')}: {subject_type}")
    feature_ids = set(feature_by_id(root))
    covered_feature_ids = {row.get("feature_id", "") for row in coverage_rows if row.get("source_kind") == "BF"}
    for feature_id in sorted(feature_ids - covered_feature_ids):
        errors.append(f"coverage.csv: BF {feature_id} не покрыт и не классифицирован")
    dirs = [subject_root(root) / "cards" / card] if card else card_dirs(root)
    if not dirs and not registry_rows:
        errors.append("Нет предметных карточек; выполните subject-card seed.")
    ready_count = 0
    card_slugs: set[str] = set()
    for card_dir in dirs:
        payload_path = card_dir / "subject-card.json"
        if not payload_path.exists():
            errors.append(f"Нет карточки: {payload_path.relative_to(root).as_posix()}")
            continue
        try:
            payload = json.loads(payload_path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"Не удалось прочитать {payload_path.relative_to(root).as_posix()}: {exc}")
            continue
        evidence_rows = non_empty_rows(read_csv_rows(card_dir / "evidence.csv"))
        gaps = non_empty_rows(read_csv_rows(card_dir / "gaps.csv"))
        if payload.get("schema_version") != "subject-card/v1":
            errors.append(f"{payload_path.relative_to(root).as_posix()}: неверный schema_version")
        for field in ("slug", "title", "status", "confidence", "summary", "identification", "key_conclusion", "upgrade_risk"):
            if not str(payload.get(field) or "").strip():
                errors.append(f"{payload.get('slug') or card_dir.name}: не заполнено поле {field}")
        slug = str(payload.get("slug") or card_dir.name).strip()
        if slug:
            card_slugs.add(slug)
        for field in ("subject_type", "origin_layer", "coverage_scope", "why_separate_card"):
            if not str(payload.get(field) or "").strip():
                errors.append(f"{slug or card_dir.name}: не заполнено поле {field}")
        subject_type = str(payload.get("subject_type") or "").strip()
        if subject_type and subject_type not in SUBJECT_CARD_TYPES:
            errors.append(f"{slug or card_dir.name}: недопустимый subject_type {subject_type}")
        if slug and slug not in registry_by_slug:
            errors.append(f"{slug}: карточка отсутствует в registry.csv")
        for source in split_refs(payload.get("source_artifacts", [])):
            if spreadsheet_reference(source):
                errors.append(f"{payload.get('slug') or card_dir.name}: Excel не может быть source_artifact карточки: {source}")
        if not isinstance(payload.get("sections"), dict):
            errors.append(f"{payload.get('slug') or card_dir.name}: sections должен быть объектом")
        else:
            for section, rows in payload["sections"].items():
                if section == "open_questions":
                    continue
                if not isinstance(rows, list):
                    errors.append(f"{payload.get('slug') or card_dir.name}: секция {section} должна быть списком")
                    continue
                for index, row in enumerate(rows, start=1):
                    if isinstance(row, dict) and not str(row.get("source") or row.get("source_path") or "").strip():
                        errors.append(f"{payload.get('slug') or card_dir.name}: строка {index} секции {section} не имеет источника")
                    if isinstance(row, dict):
                        for key in ("source", "source_path"):
                            source = str(row.get(key) or "").strip()
                            if spreadsheet_reference(source):
                                errors.append(f"{payload.get('slug') or card_dir.name}: строка {index} секции {section} ссылается на Excel как источник: {source}")
        if not evidence_rows:
            errors.append(f"{payload.get('slug') or card_dir.name}: нет evidence.csv строк")
        for evidence in evidence_rows:
            source = str(evidence.get("source_path") or "").strip()
            if spreadsheet_reference(source):
                errors.append(f"{payload.get('slug') or card_dir.name}: evidence {evidence.get('evidence_id')} ссылается на Excel как источник: {source}")
        blocking_gaps = [row for row in gaps if row.get("blocking") == "true" and row.get("status") != "closed"]
        if payload.get("status") == "ready_for_review":
            ready_count += 1
            contour = accepted_contour_by_slug.get(slug, {})
            registry = registry_by_slug.get(slug, {})
            if not contour:
                errors.append(f"{slug}: ready_for_review без accepted contour в contours.csv")
            if generic_bucket_title(str(payload.get("title") or "")) and "не техническая корзина" not in str(payload.get("why_not_technical_bucket") or "").lower():
                errors.append(f"{slug}: generic-название ready-карточки требует why_not_technical_bucket")
            for field in ("why_separate_card", "key_conclusion", "upgrade_risk"):
                if generic_template_text(str(payload.get(field) or "")):
                    errors.append(f"{slug}: поле {field} похоже на шаблонный текст")
            if not str(payload.get("migration_boundary") or "").strip():
                errors.append(f"{slug}: ready-карточка не содержит migration_boundary")
            core_objects = split_refs(payload.get("primary_objects", [])) or split_refs(contour.get("primary_objects", "")) or split_refs(registry.get("primary_objects", ""))
            if not core_objects:
                errors.append(f"{slug}: ready-карточка не содержит ядро сценария в primary_objects или accepted contour")
            if contour and registry:
                if split_refs(contour.get("primary_objects", "")) != split_refs(registry.get("primary_objects", "")):
                    errors.append(f"{slug}: primary_objects в registry.csv и accepted contour не совпадают")
                if contour.get("title", "").strip() and registry.get("title", "").strip() and contour.get("title") != registry.get("title"):
                    errors.append(f"{slug}: title в registry.csv и accepted contour не совпадает")
            if blocking_gaps:
                errors.append(f"{payload.get('slug')}: статус ready_for_review при открытых blocking gaps")
    missing_cards = [slug for slug in card_slugs if not registry_by_slug.get(slug, {}).get("card_path")]
    for slug in missing_cards:
        errors.append(f"{slug}: registry.csv не содержит card_path для существующей карточки")
    return {
        "status": "ok" if not errors else "fail",
        "ready_count": ready_count,
        "registry_count": len(registry_rows),
        "coverage_count": len(coverage_rows),
        "errors": errors,
    }


def sheet_rows(workbook_path: Path, sheet: str, keys: list[str]) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[sheet]
    rows: list[dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        row: dict[str, str] = {}
        for index, key in enumerate(keys):
            value = values[index] if index < len(values) else ""
            if value is None:
                row[key] = ""
            elif isinstance(value, float) and value.is_integer():
                row[key] = str(int(value))
            else:
                row[key] = str(value).strip()
        rows.append(row)
    return rows


def compare_reference(root: Path, card: str, workbook: Path) -> dict[str, Any]:
    payload, _, _ = load_subject_card(root, card)
    sections = payload.get("sections", {})
    mapping = [
        ("Реквизиты", "attributes", ["object", "kind", "name", "synonym", "data_type", "vendor_status", "relation", "confidence", "source", "line", "comment"]),
        ("Правила формы", "form_rules", ["id", "rule", "description", "mechanism", "confidence", "source", "line"]),
        ("Проверки заполнения", "validations", ["id", "field", "validation", "mechanism", "confidence", "source", "line"]),
        ("Жизненный цикл", "lifecycle", ["id", "action", "description", "mechanism", "confidence", "source", "line"]),
        ("Права и роли", "rights", ["id", "role", "description", "mechanism", "confidence", "source", "line"]),
        ("Источники", "sources", ["id", "claim", "source", "line", "evidence"]),
        ("Открытые вопросы", "open_questions", ["id", "question", "why_open", "needed"]),
    ]
    mismatches: list[str] = []
    for sheet, section, keys in mapping:
        excel_rows = sheet_rows(workbook, sheet, keys)
        card_rows = [{key: str(row.get(key, "")).strip() for key in keys} for row in sections.get(section, [])]
        if excel_rows != card_rows:
            mismatches.append(f"{sheet}: excel={len(excel_rows)} card={len(card_rows)}")
    return {"status": "ok" if not mismatches else "fail", "mismatches": mismatches}


def discover_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = discover_subject_cards(root)
    print(f"subject_card_candidates: {result['candidates']}")
    print(f"accepted_candidates: {result['accepted']}")
    return 0


def contour_draft_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = draft_subject_card_contours(root)
    print(f"subject_card_contours: {result['contours']}")
    print(f"accepted_contours: {result['accepted']}")
    return 0


def contour_validate_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = validate_subject_card_contours(root)
    print(f"subject_card_contour_status: {result['status']}")
    print(f"contours: {result['contours']}")
    print(f"accepted_contours: {result['accepted']}")
    for error in result["errors"]:
        print(f"- {error}")
    return 0 if result["status"] == "ok" else 1


def classify_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = classify_subject_cards(root)
    print(f"subject_card_classified: {result['classified']}")
    print(f"accepted_decisions: {result['accepted']}")
    return 0


def registry_build_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = build_subject_card_registry(root)
    print(f"subject_card_registry_rows: {result['registry_rows']}")
    print(f"subject_card_coverage_rows: {result['coverage_rows']}")
    print(f"subject_cards: {result['cards']}")
    return 0


def seed_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = seed_subject_cards(root, all_candidates=args.all, card=args.card or "", from_registry=args.from_registry)
    print(f"subject_cards_written: {len(result['written'])}")
    for slug in result["written"]:
        print(f"- {slug}")
    if result["skipped"]:
        print(f"subject_candidates_skipped: {len(result['skipped'])}")
    return 0


def refine_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = refine_subject_card(root, args.card)
    print(f"subject_card: {result['card']}")
    print(f"status: {result['status']}")
    print(f"evidence_count: {result['evidence_count']}")
    print(f"gap_count: {result['gap_count']}")
    return 0


def validate_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = validate_subject_cards(root, card=args.card or "")
    print(f"subject_card_status: {result['status']}")
    print(f"ready_for_review: {result['ready_count']}")
    print(f"registry_rows: {result.get('registry_count', 0)}")
    print(f"coverage_rows: {result.get('coverage_count', 0)}")
    for error in result["errors"]:
        print(f"- {error}")
    return 0 if result["status"] == "ok" else 1


def compare_reference_command(args: argparse.Namespace) -> int:
    root = Path(args.repo_path).resolve() if args.repo_path else Path.cwd()
    result = compare_reference(root, args.card, Path(args.workbook).resolve())
    print(f"subject_card_reference_status: {result['status']}")
    for mismatch in result["mismatches"]:
        print(f"- {mismatch}")
    return 0 if result["status"] == "ok" else 1
